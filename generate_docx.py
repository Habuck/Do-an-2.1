import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

def main():
    doc = Document()
    doc.add_heading('Toàn bộ Source Code Quan Trọng - Project Easy Kit Audio', 0)
    
    files_info = [
        {
            "path": "web_app.py",
            "desc": "File chính của backend sử dụng Flask. Chứa các API endpoints để giao tiếp giữa giao diện web (frontend) và các mô hình xử lý âm thanh, AI (backend). Xử lý CORS và routing cơ bản."
        },
        {
            "path": "index.html",
            "desc": "Giao diện chính của ứng dụng web (Frontend). Bao gồm giao diện người dùng, thu âm giọng nói, và gọi API tới server."
        },
        {
            "path": "audio_project/core.py",
            "desc": "Chứa các hàm cốt lõi để xử lý âm thanh cơ bản như trích xuất đặc trưng (MFCC, pitch, năng lượng), chuẩn hóa và các tiền xử lý trước khi đưa vào mô hình AI."
        },
        {
            "path": "audio_project/emotion.py",
            "desc": "Mô-đun nhận diện cảm xúc giọng nói. Có thể tích hợp mô hình pre-trained wav2vec2 hoặc mô hình tương tự để phân loại cảm xúc (vui, buồn, tức giận, bình thường)."
        },
        {
            "path": "audio_project/gender.py",
            "desc": "Mô-đun dự đoán giới tính (nam/nữ) dựa trên giọng nói. Thường sử dụng các thuật toán như K-Nearest Neighbors (KNN) kết hợp với các đặc trưng âm thanh được trích xuất."
        },
        {
            "path": "audio_project/pronunciation.py",
            "desc": "Mô-đun đánh giá phát âm. Sử dụng thuật toán DTW (Dynamic Time Warping) để so sánh audio thu âm với audio mẫu, từ đó chấm điểm phát âm tiếng Việt."
        },
        {
            "path": "audio_project/cough.py",
            "desc": "Mô-đun phân tích tiếng ho, ứng dụng trong việc chẩn đoán hoặc phát hiện các đặc trưng liên quan đến sức khỏe qua tiếng ho."
        },
        {
            "path": "audio_project/depression.py",
            "desc": "Mô-đun phân tích trầm cảm qua giọng nói. Sử dụng AI để tìm kiếm các dấu hiệu bất thường trong âm điệu biểu hiện trầm cảm."
        }
    ]

    for item in files_info:
        file_path = item["path"]
        desc = item["desc"]
        
        if not os.path.exists(file_path):
            continue
            
        doc.add_heading(file_path, level=1)
        doc.add_paragraph(desc, style='Intense Quote')
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except Exception as e:
            content = f"Lỗi không thể đọc file: {e}"
            
        # Thêm code với font monospace
        p = doc.add_paragraph()
        run = p.add_run(content)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        
        doc.add_page_break()

    doc.save('SourceCode_QuanTrong.docx')
    print("Đã tạo file SourceCode_QuanTrong.docx thành công!")

if __name__ == "__main__":
    main()
